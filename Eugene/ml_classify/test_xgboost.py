"""
Test XGBoost classifier on 1k jobs
Compare with JoBert and Random Forest results
"""
import pandas as pd
import joblib
from collections import defaultdict
from tqdm import tqdm
import os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from sentence_splitter_v2 import split_into_sentences_v2

# ============================================================================
# Configuration
# ============================================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_PATH = os.path.dirname(SCRIPT_DIR)

INPUT_FILE = os.path.join(BASE_PATH, "diverse_sample_1k.parquet")
OUTPUT_EXCEL = os.path.join(SCRIPT_DIR, "categorized_jobs_xgb.xlsx")

MODEL_PATH = os.path.join(SCRIPT_DIR, "xgb_classifier.joblib")
VECTORIZER_PATH = os.path.join(SCRIPT_DIR, "tfidf_vectorizer_xgb.joblib")
ENCODER_PATH = os.path.join(SCRIPT_DIR, "label_encoder.joblib")

# ============================================================================
# Load Model
# ============================================================================
print("="*80)
print("XGBOOST PIPELINE")
print("="*80)

print("\nLoading model...")
xgb_clf = joblib.load(MODEL_PATH)
vectorizer = joblib.load(VECTORIZER_PATH)
label_encoder = joblib.load(ENCODER_PATH)
print("Model loaded!")

label_names = list(label_encoder.classes_)
print(f"Labels: {label_names}")

# ============================================================================
# Load Data
# ============================================================================
print("\nLoading data...")
df = pd.read_parquet(INPUT_FILE)
df_sample = df[["job_id", "job_description"]].copy()
print(f"Processing {len(df_sample)} jobs")

# ============================================================================
# Process Jobs
# ============================================================================
print("\n" + "="*80)
print("Processing jobs...")
print("="*80)

results = []
stats = {'total': 0, 'classified': 0}

for idx, row in tqdm(df_sample.iterrows(), total=len(df_sample)):
    job_id = row['job_id']
    job_desc = str(row['job_description'])

    # Split into sentences using V2 splitter
    sentences = split_into_sentences_v2(job_desc)

    label_sentences = defaultdict(list)

    for sentence in sentences:
        if not sentence.strip():
            continue

        stats['total'] += 1

        # Classify with XGBoost
        X = vectorizer.transform([sentence])
        prediction_idx = xgb_clf.predict(X)[0]
        prediction = label_encoder.inverse_transform([prediction_idx])[0]

        if prediction in label_names:
            label_sentences[prediction].append(sentence)
            stats['classified'] += 1

    results.append({
        'job_id': job_id,
        'job_description': job_desc,
        'About the Company': " ".join(label_sentences.get('About the Company', [])),
        'Job Description': " ".join(label_sentences.get('Job Description', [])),
        'Job Requirements': " ".join(label_sentences.get('Job Requirements', [])),
        'Responsibilities': " ".join(label_sentences.get('Responsibilities', [])),
        'Benefits': " ".join(label_sentences.get('Benefits', [])),
        'sentence_count': len(sentences)
    })

# ============================================================================
# Save Results
# ============================================================================
print("\n" + "="*80)
print("Saving results...")
print("="*80)

results_df = pd.DataFrame(results)
results_df.to_excel(OUTPUT_EXCEL, index=False)

# Format the Excel file for readability
print("Formatting Excel file...")
wb = load_workbook(OUTPUT_EXCEL)
ws = wb.active

# Define styles
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
wrap_alignment = Alignment(wrap_text=True, vertical="top")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Format header row
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Format data cells
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = wrap_alignment
        cell.border = thin_border

# Set column widths
column_widths = {
    'A': 15,   # job_id
    'B': 60,   # job_description
    'C': 40,   # About the Company
    'D': 50,   # Job Description
    'E': 50,   # Job Requirements
    'F': 50,   # Responsibilities
    'G': 40,   # Benefits
    'H': 12,   # sentence_count
}
for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

# Freeze header row
ws.freeze_panes = 'A2'

wb.save(OUTPUT_EXCEL)
print(f"Saved to: {OUTPUT_EXCEL}")

# ============================================================================
# Summary
# ============================================================================
print("\n" + "="*80)
print("SUMMARY")
print("="*80)
print(f"Total sentences: {stats['total']:,}")
print(f"Classified: {stats['classified']:,}")

print("\n### Category Coverage ###")
for cat in label_names:
    has_content = (results_df[cat].notna() & (results_df[cat] != '')).sum()
    print(f"  {cat:20}: {has_content:4} jobs ({100*has_content/len(results_df):5.1f}%)")

print("\nDone!")
