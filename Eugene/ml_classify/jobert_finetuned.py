"""
JoBert Fine-tuned Pipeline
- Uses fine-tuned JoBert model (91% accuracy vs ~60% original)
- Includes PII masking (emails, phones, URLs, CEI/EA numbers)
- Filters contact/application sentences
- Filters section headers
"""
import torch
from transformers import AutoTokenizer, AutoModelForSequenceClassification
import pandas as pd
import re
import html
from collections import defaultdict
from tqdm import tqdm
from jobert_utils import preprocess_line_breaks, split_into_sentences
tqdm.pandas()

# ============================================================================
# Configuration
# ============================================================================
import os

# Get the directory where this script is located
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
# Eugene folder is one level up from ml_classify
BASE_PATH = os.path.dirname(SCRIPT_DIR)

USE_FINETUNED_MODEL = True  # Set to False to use original AhmedBou/JoBert
FINETUNED_MODEL_PATH = os.path.join(BASE_PATH, "jobert_finetuned")

# Input/Output paths (relative to Eugene folder)
INPUT_FILE = os.path.join(BASE_PATH, "sampled_jobs_final.parquet")  # Change this to your input file
OUTPUT_PARQUET = os.path.join(SCRIPT_DIR, "categorized_jobs_finetuned.parquet")
OUTPUT_EXCEL = os.path.join(SCRIPT_DIR, "categorized_jobs_finetuned.xlsx")

# Sample size (set to None to process all)
SAMPLE_SIZE = None  # Process all jobs in the input file

# ============================================================================
# PII Masking Functions
# ============================================================================
def mask_pii(text: str) -> str:
    """
    Mask PII (emails, phones, URLs) with ****
    Based on Eugene's pii.py mask_contacts function
    """
    if not text or not isinstance(text, str):
        return text

    # Mask emails
    text = re.sub(r'\b[\w\.-]+@[\w\.-]+\.\w+\b', '****', text)

    # Mask phone numbers (various formats)
    text = re.sub(r'\b(?:\+?\d{1,3}[ -]?)?(?:\d[ -]?){6,14}\b', '****', text)

    # Mask URLs and domains
    text = re.sub(
        r'(?:https?://|www\.|[a-zA-Z0-9.-]+\.(?:com|sg|net|org|io|edu|gov))\S*',
        '****',
        text,
        flags=re.IGNORECASE
    )

    # Mask CEI/EA registration numbers (Singapore recruitment)
    text = re.sub(r'\b[A-Z]{1,2}\d{5,8}[A-Z]?\b', '****', text)
    text = re.sub(r'\bEA\s*(?:License|Licence|Reg|Registration)?\s*(?:No\.?|Number)?:?\s*\d+[A-Z]*\b', '****', text, flags=re.IGNORECASE)
    text = re.sub(r'\bCEI\s*(?:Reg|Registration)?\s*(?:No\.?|Number)?:?\s*[A-Z]*\d+[A-Z]*\b', '****', text, flags=re.IGNORECASE)

    return text


def is_contact_sentence(sentence: str) -> bool:
    """
    Check if sentence is primarily contact/application info that should be filtered
    """
    sentence_lower = sentence.lower()

    contact_keywords = [
        'email your resume', 'email us', 'send your cv', 'send your resume',
        'apply now', 'apply at', 'apply to', 'apply via',
        'contact us', 'contact:', 'reach us',
        'regret that only', 'shortlisted candidates', 'short-listed',
        'we regret', 'only shortlisted',
        'cei registration', 'ea license', 'ea licence', 'personnel no',
        'registration no', 'license no', 'licence no',
    ]

    return any(kw in sentence_lower for kw in contact_keywords)


def is_section_header(sentence: str) -> bool:
    """
    Check if sentence is just a section header like 'Responsibilities:' that should be filtered
    """
    cleaned = sentence.strip().rstrip(':').lower()

    section_headers = {
        'responsibilities', 'responsibility', 'job responsibilities', 'key responsibilities',
        'requirements', 'requirement', 'job requirements', 'key requirements',
        'qualifications', 'qualification', 'qualifications & skills',
        'about the company', 'about us', 'company overview', 'who we are',
        'job description', 'job summary', 'role overview', 'the role',
        'benefits', 'what we offer', 'perks', 'compensation',
        'duties', 'key duties', 'job duties',
        'skills', 'skills required', 'required skills',
        'experience', 'education', 'overview', 'summary', 'description',
    }

    return cleaned in section_headers


# ============================================================================
# Load Data
# ============================================================================
print("="*80)
print("JOBERT FINE-TUNED PIPELINE")
print("="*80)

df = pd.read_parquet(INPUT_FILE)
print(f"Original dataframe shape: {df.shape}")
print(f"Columns: {list(df.columns)}")

# Sample rows
if SAMPLE_SIZE:
    df_sample = df.sample(n=min(SAMPLE_SIZE, len(df)), random_state=42)
else:
    df_sample = df.copy()
df_sample = df_sample[["job_id", "job_description"]].copy()
print(f"Processing {len(df_sample)} jobs")

# ============================================================================
# Load Model
# ============================================================================
print("\n" + "="*80)
print("Loading model...")
print("="*80)

if USE_FINETUNED_MODEL:
    print(f"Using FINE-TUNED model from: {FINETUNED_MODEL_PATH}")
    tokenizer = AutoTokenizer.from_pretrained(FINETUNED_MODEL_PATH)
    model = AutoModelForSequenceClassification.from_pretrained(FINETUNED_MODEL_PATH)
else:
    print("Using ORIGINAL AhmedBou/JoBert model")
    tokenizer = AutoTokenizer.from_pretrained("AhmedBou/JoBert")
    model = AutoModelForSequenceClassification.from_pretrained("AhmedBou/JoBert")

model.eval()

# Setup device
if torch.backends.mps.is_available():
    device = torch.device("mps")
    print("Using Apple Silicon GPU (MPS)")
elif torch.cuda.is_available():
    device = torch.device("cuda")
    print("Using CUDA GPU")
else:
    device = torch.device("cpu")
    print("Using CPU")

model = model.to(device)

label_names = ['About the Company', 'Job Description', 'Job Requirements', 'Responsibilities', 'Benefits']


def classify_sentence(sentence, model, tokenizer, label_names):
    """Classify a single sentence"""
    try:
        inputs = tokenizer(sentence, return_tensors='pt', truncation=True, max_length=128)
        inputs = {k: v.to(device) for k, v in inputs.items()}

        with torch.no_grad():
            outputs = model(**inputs)

        prediction = torch.argmax(outputs.logits).item()

        if prediction < len(label_names):
            return label_names[prediction]
        return None
    except Exception as e:
        print(f"Error classifying sentence: {e}")
        return None


# ============================================================================
# Process Jobs
# ============================================================================
print("\n" + "="*80)
print("Processing job descriptions...")
print("="*80)

# Preprocess job descriptions
print("Step 1: Preprocessing (HTML cleaning)...")
df_sample['job_description_clean'] = df_sample['job_description'].progress_apply(preprocess_line_breaks)

# Mask PII
print("Step 2: Masking PII...")
df_sample['job_description_masked'] = df_sample['job_description_clean'].progress_apply(mask_pii)

# Initialize columns
for label in label_names:
    df_sample[label] = ""
df_sample['sentences'] = None
df_sample['sentence_count'] = 0

# Stats tracking
stats = {'total_sentences': 0, 'filtered_contact': 0, 'filtered_headers': 0, 'classified': 0}

# Process each row
print("Step 3: Classifying sentences...")
for idx, row in tqdm(df_sample.iterrows(), total=len(df_sample)):
    job_desc = row['job_description_masked']
    sentences = split_into_sentences(job_desc, smart_split=True)

    df_sample.at[idx, 'sentences'] = sentences
    df_sample.at[idx, 'sentence_count'] = len(sentences)

    label_sentences = defaultdict(list)

    for sentence in sentences:
        if not sentence.strip():
            continue

        stats['total_sentences'] += 1

        # Filter contact sentences
        if is_contact_sentence(sentence):
            stats['filtered_contact'] += 1
            continue

        # Filter section headers
        if is_section_header(sentence):
            stats['filtered_headers'] += 1
            continue

        # Classify
        label = classify_sentence(sentence, model, tokenizer, label_names)
        if label:
            label_sentences[label].append(sentence)
            stats['classified'] += 1

    # Assign categorized sentences to columns
    for label in label_names:
        df_sample.at[idx, label] = " ".join(label_sentences[label])

# Drop intermediate columns
df_sample = df_sample.drop(columns=['job_description_clean', 'job_description_masked'])

print("\nClassification complete!")

# ============================================================================
# Save Results
# ============================================================================
print("\n" + "="*80)
print("Saving results...")
print("="*80)

# Save to parquet
df_sample.to_parquet(OUTPUT_PARQUET, index=False)
print(f"Saved to parquet: {OUTPUT_PARQUET}")

# Clean for Excel
def clean_for_excel(text):
    """Remove control characters and unescape HTML"""
    if isinstance(text, list):
        return [clean_for_excel(item) for item in text]
    if not isinstance(text, str):
        return text
    text = html.unescape(text)
    text = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F]', '', text)
    return text

df_excel = df_sample.copy()
for col in df_excel.columns:
    if df_excel[col].dtype == 'object':
        df_excel[col] = df_excel[col].apply(clean_for_excel)

# Save with formatting
from openpyxl.styles import Alignment, Font, PatternFill

writer = pd.ExcelWriter(OUTPUT_EXCEL, engine='openpyxl')
df_excel.to_excel(writer, index=False, sheet_name='Results')

workbook = writer.book
worksheet = writer.sheets['Results']

# Set column widths
column_widths = {'A': 40, 'B': 80, 'C': 60, 'D': 60, 'E': 60, 'F': 60, 'G': 60, 'H': 60, 'I': 15}
for col, width in column_widths.items():
    worksheet.column_dimensions[col].width = width

# Style header row
header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF')
for cell in worksheet[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Style data cells
for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
    for cell in row:
        cell.alignment = Alignment(vertical='top', wrap_text=True)

worksheet.row_dimensions[1].height = 30
worksheet.freeze_panes = 'A2'
writer.close()

print(f"Saved to Excel: {OUTPUT_EXCEL}")

# ============================================================================
# Summary
# ============================================================================
print("\n" + "="*80)
print("SUMMARY")
print("="*80)
print(f"\nTotal jobs processed: {len(df_sample)}")
print(f"Total sentences: {stats['total_sentences']:,}")
print(f"Contact sentences filtered: {stats['filtered_contact']:,}")
print(f"Section headers filtered: {stats['filtered_headers']:,}")
print(f"Sentences classified: {stats['classified']:,}")

print("\n### Category Coverage ###")
for label in label_names:
    has_content = (df_sample[label].notna() & (df_sample[label] != '')).sum()
    print(f"  {label:20}: {has_content:4} jobs ({100*has_content/len(df_sample):5.1f}%)")

print(f"\nColumns: {list(df_sample.columns)}")
print("\n" + "="*80)
print("Done!")
print("="*80)
